import json
import shutil
from pathlib import Path
from datetime import datetime, timezone
from copy import deepcopy
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[2]
ORIGINAL = ROOT / '全球黄金迁徙地图上线需求.docx'
OUT = ROOT / '全球黄金迁徙地图上线需求.docx'
d = Document(ORIGINAL) if ORIGINAL.exists() else Document()
for item in list(d._element.body):
    if item.tag != qn('w:sectPr'):
        d._element.body.remove(item)
for rel in list(d.part.rels.values()):
    if rel.reltype == RT.HYPERLINK:
        d.part.drop_rel(rel.rId)

font_name = 'Arial Unicode MS'
for name, size in [('Normal',10.5),('Title',18),('Heading 1',13),('Heading 2',11.5)]:
    st = d.styles[name]
    st.font.name = font_name
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor(0,0,0)
    st.font.underline = False
    st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),font_name)
    st._element.get_or_add_rPr().rFonts.set(qn('w:ascii'),font_name)
    st._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'),font_name)
    for border in list(st._element.xpath('./w:pPr/w:pBdr')):
        border.getparent().remove(border)
    f = st.paragraph_format
    f.line_spacing = 1.15
    f.space_after = Pt(5)
    f.space_before = Pt(0 if name=='Normal' else 9)
    f.keep_with_next = name != 'Normal'
    f.widow_control = True
d.styles['Title'].paragraph_format.space_after = Pt(8)
for s in d.sections:
    s.page_width,s.page_height = Cm(21),Cm(29.7)
    s.left_margin,s.right_margin = Cm(2),Cm(2)
    s.top_margin,s.bottom_margin = Cm(1.8),Cm(1.8)
    s.header_distance,s.footer_distance = Cm(.7),Cm(.8)
    for p in s.header.paragraphs:p.clear()
    for p in s.footer.paragraphs:p.clear()
    p=s.footer.paragraphs[0]
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run();r.font.size=Pt(8)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE')
    r._r.addnext(field)
d.core_properties.title = '全球黄金迁徙地图上线需求'
d.core_properties.subject = '功能 数据来源 趋势判断 抓取更新与上线验收'
d.core_properties.author = ''
d.core_properties.last_modified_by = ''
d.core_properties.comments = ''
d.core_properties.modified = datetime.now(timezone.utc)

def p(text='', bold_lead=None):
    x=d.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        x.add_run(bold_lead).bold=True;x.add_run(text[len(bold_lead):])
    else:x.add_run(text)
    return x

def h(text, level=1):
    return d.add_paragraph(text,style=f'Heading {level}')

def link(x,label,url):
    rid=d.part.relate_to(url,RT.HYPERLINK,is_external=True)
    hl=OxmlElement('w:hyperlink');hl.set(qn('r:id'),rid)
    run=OxmlElement('w:r');rp=OxmlElement('w:rPr')
    fonts=OxmlElement('w:rFonts')
    for key in ['ascii','hAnsi','eastAsia']:fonts.set(qn('w:'+key),font_name)
    rp.append(fonts)
    color=OxmlElement('w:color');color.set(qn('w:val'),'155E82');rp.append(color)
    u=OxmlElement('w:u');u.set(qn('w:val'),'single');rp.append(u)
    run.append(rp);t=OxmlElement('w:t');t.text=label;run.append(t);hl.append(run);x._p.append(hl)

def links(label,items):
    x=p(label)
    for i,(title,url) in enumerate(items):
        if i:x.add_run('；')
        link(x,title,url)
    return x

def table(headers,rows,widths):
    t=d.add_table(rows=1, cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for col,width in zip(t.columns,widths):col.width=Cm(width)
    props=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    props.append(borders)
    margin=OxmlElement('w:tblCellMar')
    for edge,val in [('top','65'),('bottom','65'),('left','85'),('right','85')]:
        e=OxmlElement('w:'+edge);e.set(qn('w:w'),val);e.set(qn('w:type'),'dxa');margin.append(e)
    props.append(margin)
    rep=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(rep)
    for ri,vals in enumerate([headers]+rows):
        row=t.rows[0] if ri==0 else t.add_row()
        nosplit=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(nosplit)
        for ci,(cell,text,width) in enumerate(zip(row.cells,vals,widths)):
            cell.width=Cm(width);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text=text
            if ri==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E5EFF5');cell._tc.get_or_add_tcPr().append(shade)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(2)
                para.paragraph_format.space_before=Pt(0)
                para.paragraph_format.line_spacing=1.1
                para.paragraph_format.keep_with_next=False
                for r in para.runs:
                    r.font.size=Pt(9.5);r.bold=ri==0;r.font.color.rgb=RGBColor(0,0,0)
    p('').paragraph_format.space_after=Pt(1)
    return t


POLICY = json.loads((ROOT / 'gold-migration-demo/app/data/source-policy.json').read_text())
S = POLICY['sources']
M = {x['code']:x for x in POLICY['markets']}

# Both this document and the demo derive source identities from this registry.
def trade_table(codes, purpose):
    p('抓取内容：'+purpose+' 商品取非货币黄金710811、710812、710813及对应本地细分码，重量统一为吨，保留月份、数量单位、金额币种、估算和修订状态。当前Comtrade HS7108不报告710820货币黄金，不覆盖首饰或废料等全部黄金。')
    t=table(['市场','优先接入主源','demo当前执行与接入条件','抓取频次'],[],[1.6,3.7,7.2,4.5])
    for code in codes:
        r=M[code];chosen=S[r['selectedSourceId']];current=S[r['currentSourceId']]
        row=t.add_row()
        vals=[r['market'],chosen['name'],current['name']+'。'+r['gap'],'每日09:00；明确发布日起每6小时查48小时。']
        for cell,text,width in zip(row.cells,vals,[1.6,3.7,7.2,4.5]):
            cell.width=Cm(width);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;cell.text=text
            for para in cell.paragraphs:
                para.paragraph_format.keep_with_next=False
                para.paragraph_format.space_after=Pt(3)
                para.paragraph_format.line_spacing=1.1
                for run in para.runs:run.font.size=Pt(9.5)
        row.cells[1].text='';link(row.cells[1].paragraphs[0],chosen['name'],chosen['url'])
        row.cells[2].text='';x=row.cells[2].paragraphs[0]
        link(x,current['name'],current['url']);x.add_run('。'+r['gap'])
        no=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no)
    p('来源按月发布，各市场时间不同。每次回查近3个月，每月回查18个月修订。每个市场只有一个当前执行源；表中优先源通过验收后，统一替换该市场所有依赖板块及历史比较期。未通过前继续使用当前执行源，并标真实观测月，不跨源补缺。')
    if 842 in codes:p('美国Census数据免费，API需申请并激活免费密钥。其他待接入官方门户的完整黄金重量和导出方式须实取确认；公开入口不等于已经可抓取完整明细。')
    if 702 in codes:p('新加坡使用免费SingStat作为优先验证渠道；StatLink将于2027年1月1日停用，不新增长期依赖。')

def metric_table(keys):
    t=table(['数据与指定来源','抓取内容和访问条件','抓取频次'],[],[4.3,7.2,5.5])
    for key in keys:
        r=POLICY['metrics'][key];chosen=S[r['selectedSourceId']];current=S[r['currentSourceId']]
        row=t.add_row()
        vals=[r['label'],r['scope']+' '+chosen['releaseCycle']+'。'+chosen['access']+' '+r['gap'],r['checkSchedule']]
        for cell,text,width in zip(row.cells,vals,[4.3,7.2,5.5]):
            cell.width=Cm(width);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;cell.text=text
            for para in cell.paragraphs:
                para.paragraph_format.keep_with_next=False;para.paragraph_format.space_after=Pt(3);para.paragraph_format.line_spacing=1.1
                for run in para.runs:run.font.size=Pt(9.5)
        cell=row.cells[0];cell.text='';x=cell.paragraphs[0];x.add_run(r['label']).bold=True
        x=cell.add_paragraph();link(x,chosen['name'],chosen['url'])
        if chosen['fetchUrl']:
            x=cell.add_paragraph();link(x,'实际公开取数地址',chosen['fetchUrl'])
        if r['currentSourceId']!=r['selectedSourceId']:
            x=cell.add_paragraph('当前demo：'+current['name'])
        no=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no)
    p('')

def lead(label,text):return p(label+'：'+text,bold_lead=label+'：')
def ai(text):lead('可选AI提示词',text+' 只润色已确认结论，不计算、补数据或编造因果。数字另行展示，地区名称统一为中国台湾。')
def module(title):
    heading = h(title)
    if title.startswith('06 '):
        heading.paragraph_format.page_break_before = True

d.add_paragraph('全球黄金迁徙地图上线需求',style='Title')
p('交付对象 技术团队    日期 2026年10月10日')
p('按照当前demo的六个板块上线可持续更新的网站。每个板块在本节说明要做什么、抓取什么、来源、频次、趋势和验收。主句先展示由西向东、东向放缓、库存增加等判断，数字作为依据。')
links('参考页面 ',[('demo','https://global-gold-migration-demo.kinardqgumt.chatgpt.site'),('GitHub仓库','https://github.com/JennyChiu68/gold-migration')])
p('来源方案日期 '+POLICY['revision']+'。优先接入主源与demo当前执行源分别写明；尚未通过完整明细、重量、伙伴编码、重叠月份和访问条件验收的源，不写成已经使用。多个板块复用同一份数据，换源后一起更新。')
p('免费可获取标准：本阶段只要无需付费、能从公开入口实际抓取，即视为免费可获取；版权和许可不作为来源筛选或接入前置条件。免费注册、密钥、限流与实际抓取失败仍需注明。美国Census密钥可免费申请，BAZG公开CSV可直接下载。伦敦PM公开JSON当前返回403，保留检查和重试；不将付费交付服务列入当前方案。')
p('当前demo是固定快照，尚无自动定时任务。下文频次是技术需要实现的安排，默认北京时间，注明当地时间的按当地工作日和夏令时执行。各卡片标真实观测期；历史数据不参与当前趋势。当前demo未接AI，上线不强制采用；可选AI仅润色规则先得出的结论。页面、列表、地图、筛选和文案统一使用中国台湾。')

module('01 黄金罗盘')
lead('要做什么','用报关、库存、现货三条短句概括监测方向。保留三张数字卡片与状态标签；伦敦价格尚未更新时，现货标历史观察和当前暂不可判断。')
lead('报关来源与抓取','当前六个枢纽均执行UN Comtrade。取完整月度出口目的地重量与世界总量，覆盖瑞士、英国、美国、中国香港、新加坡、阿联酋。按下列一一对应的优先源接入，不重复抓取第03板块的数据。')
x=p('优先源对应 ')
for n,c in enumerate([757,826,842,344,702,784]):
    if n:x.add_run('；')
    r=M[c];link(x,S[r['selectedSourceId']]['name'],S[r['selectedSourceId']]['url'])
links('当前报关取数 ',[('UN Comtrade取数说明',S['comtrade']['url']),('公开API入口',S['comtrade']['fetchUrl'])])
p('六个优先源均待完整性或访问验收，当前仍为Comtrade。每日09:00检查；明确发布日起每6小时查48小时，之后恢复每日；每次回查近3个月，每月回查18个月修订。各源月度更新，具体日看发布目录。美国密钥免费申请；新加坡不新增StatLink依赖。')
metric_table(['london-vault','shanghai-price','london-price','fx'])
lead('趋势判断','西方枢纽瑞士、英国、美国与东方枢纽中国香港、新加坡、阿联酋按项目固定目的地分组汇总。两方向同月完整可比，东向大于西向写监测报关流以西向东为主；反之写以东向西为主。相同枢纽和目的地范围较上月减少写东向放缓，增加写东向增强；东向仍占优但减少，可写东向仍占优但流量放缓。库存比较相邻月份；价格只用同日齐全的三项输入。')
lead('缺失与验收','覆盖减少不能写放缓；缺方向或上月时写暂无法判断。复用第03板块同一分组和全部路线，不能只加地图前40条。库存、报关、现货各有自己的日期；历史价差退出当前现货结论，其余信号继续独立更新。')
ai('根据已确认的报关方向、流量变化、库存变化及价差状态分别写短句，可写东向仍占优但流量放缓。保留各自日期，历史价差仅描述当时状态，不判断当前，不将库存或价差当作跨境迁移证明。')

module('02 全球吸金榜')
lead('要做什么','展示净流入、净流出同期榜及各市场最新完整观测。同期榜只比较同月、进口和出口重量都完整的市场；独立卡片各自标月，不跨月排名。')
trade_table([156,344,699,764,792,784,702,826,842,757],'各市场同月世界进口总量和总出口量、上月值，不用金额除金价补吨数。')
lead('趋势判断','进口减总出口，正值为净流入，负值为净流出；不足页面显示精度写基本平衡。同号比较净额绝对值，写净流入扩大或收窄、净流出扩大或收窄；跨零写转为净流入或净流出。缺上月只判断本期状态。')
lead('缺失与验收','完整新月到齐才更新该市场卡片；共同完整月推进才重算同期榜。缺重量、保密、估算、真实零分别保存；缺一侧不进榜和最新完整卡片。官方新源至少核对3个重叠月，解释差异后才启用。英国尚存在零净重附加数量非零，不能直接把缺失当零。')
ai('根据已确认的净流入或净流出及与上月变化，写一句净流入扩大、净流出收窄或转为净流入。上月不完整只写本期状态，不把净进口等同于投资需求。')

module('03 黄金航线')
lead('要做什么','保留枢纽筛选、排行榜与地图切换、路线点选和高亮。路线显示出发地、目的地、月份、吨数、原始申报额及币种、估算标记。不同官方渠道金额可能不是美元，未有可靠换算不冒充美元。')
trade_table([757,826,842,344,702,784],'报告枢纽完整出口目的地明细、世界出口总量及上月对应明细。与黄金罗盘共用一次取数。')
lead('方向分组','西方枢纽为瑞士、英国、美国；东方枢纽为中国香港、新加坡、阿联酋。东向目的地为孟加拉国、中国内地、中国香港、印尼、日本、韩国、马来西亚、其他亚洲地区（来源代码490）、巴基斯坦、印度、新加坡、越南、泰国。西向目的地为澳大利亚、奥地利、比利时、加拿大、法国、德国、意大利、荷兰、波兰、西班牙、瑞典、瑞士、英国、美国。按demo项目分组，前后期固定；代码490合并范围不当作中国台湾单独值。')
lead('趋势判断','完整目的地数据判断主要流向亚洲或某地区。路线份额用同月官方世界总量作分母；同枢纽、同范围的相邻月份份额可比才写去向重心转移。转口可能重复记录，不称全球唯一黄金流量。')
lead('缺失与验收','取齐全部分页后可展示前40条，方向计算使用全部有效路线。Comtrade预览达到500条时停止更新，拆分或取得完整交付后再发布。缺伙伴重量或世界总量，不出整体方向或份额；对方进口只核验，不能静默补入出口。未报送枢纽标覆盖范围，不能伪造当前路线。')
ai('根据已确认的主要目的地区域及相邻月份份额变化写一句路线判断。只有本期只写主要去向，可比才写重心转移；限定监测报关路线，不写全球黄金全面迁移。')

module('04 瑞士精炼站')
lead('要做什么','保留进口来源和出口去向切换、各侧地区排名及占比。进口保留全部、矿产供应地、金融及转口筛选与其余来源。两侧分别标月份和范围。')
metric_table(['swiss-import','swiss-export'])
lead('抓取补充','进口CSV按7108.1200控制代码911至914筛选，保存来源国、净重、月份、暂定状态；只保留必要月度子集，不每天保存全部大文件。出口取完整目的地与世界总量；SwissImpex通过验证前执行Comtrade。来源地角色是分析分类，不证明每批黄金的矿山产地。BAZG公开CSV直接抓取，不增加许可前置条件。')
lead('趋势与验收','完整本期写主要来自或去往某地区；前后同口径份额齐全才写来源倾斜或去向重心转移。占比用完整该侧总量。两侧范围不同，不相减推库存；修订重算对应侧排名，切换后月份和分母同步。')
ai('分别概括瑞士进口主要来源和出口主要去向；只有可比份额变化才写倾斜或转移。保留两侧日期和范围，不把来源地角色当矿山产地，不以差额推库存。')

module('05 三地实物信号')
lead('要做什么','并列伦敦月库存趋势、COMEX Registered和Eligible及总库存日变化、上海月出库与交割，保留多市场解读。')
metric_table(['london-vault','comex-vault','shanghai-physical'])
lead('报关来源与频次','多市场解读中的报关复用第03板块六枢纽数据。当前执行Comtrade；优先接入依次对应瑞士SwissImpex、英国HMRC、美国Census、中国香港Trade-IDDS、新加坡SingStat、阿联酋FCSC，均通过验证后一起切换。取完整月度出口明细及世界总量；每日09:00查，明确发布日起每6小时查48小时，校验后重算解读，无需再次下载。')
lead('趋势与验收','伦敦相邻月份和COMEX相邻活动日比较写库存增加或减少；上海月出库与上月比较写出库增加或减少，交割单列。不同方向且证据充分可写实物信号分化，保留各自指标和日期，不推断从一地运到另一地。COMEX超过两个美国营业日无可核实新表，标历史存档并退出当前解读。')
ai('分别描述已经确认的伦敦库存、COMEX库存、上海出库及监测报关变化。保留日期和指标含义，历史存档不参与当前趋势，不将不同指标相加或编造迁移因果。')

module('06 金价温差')
lead('要做什么','保留上海Au99.99对伦敦PM的指示性溢价或折价、两地价格及共同日期。当前伦敦PM仅为历史快照，页面显示历史观察和待更新，不能叫当前溢价。')
metric_table(['shanghai-price','london-price','fx'])
lead('为何需要三源','上金所提供人民币上海价格，LBMA公开JSON入口提供美元伦敦PM，ECB提供换算汇率；三个输入全部需要。伦敦公开入口当前返回403，按本节频次检查和重试；取到新数据并验证前，不推进价差日期。若发现其他无需付费的同口径公开入口，须核对日期、PM口径与历史重叠值后统一接入。采用不同价格指标须先确认定义，再回算历史，不静默替换。')
lead('趋势判断','伦敦美元每盎司价格除以31.1034768，再乘每美元人民币汇率，得到人民币每克。上海价更高为溢价、更低为折价；价差占伦敦换算价为百分比。同日三项齐全且上一共同日期也齐全才写溢价扩大或收窄、转为折价或溢价。只有一个日期只写该日期状态。')
lead('缺失与验收','取最近共同自然日期，不跨日拼接；缺项保留上一完整日期并标历史。上海收盘、伦敦PM和ECB同日也不同时，显示非实时指示性价差，不称实时套利。不能用期货价、不同定义WGC溢价或早间人民币中间价静默补缺。公开抓取成功、确认数据为最新且三源同日齐全后，才启用当前现货判断。')
ai('根据同日已确认溢价或折价及上一共同日期变化写短句，优先上海溢价扩大、溢价收窄或转为折价。历史日期仅描述当时状态，不写当前结论，不解释为可实现运输套利。')

h('交付与更新要求')
lead('来源验收','各国优先主源实取完整黄金重量、单位及伙伴表，检查分页、世界总量、金额币种、零值、估算和保密，至少3个重叠月份与现源对账并解释差异。尚未实际取到完整数据不认定完成；版权和许可不作为本阶段接入前置条件。来源方案、demo实际出处和本文对应，统一维护source-policy.json；切换须更新全部依赖板块和比较历史，不能只改网址。')
lead('自动更新','按各板块频次发现最新文件或月份，不能一直写死共同月或文件名。任一新期或修订通过校验后重算相关板块，4小时内发布并同步GitHub。没新数据不更改观测期；抓取时间、发布日、数据期分别记录。')
lead('失败处理','失败后15分钟和60分钟重试，遵守接口Retry-After；仍失败保留原期，下个周期继续。连续两轮失败通知维护人。缺上期或上期为零不计算增减百分比。月报无统一发布日期时查目录，不虚构新期；部署失败保留可用版本，确认成功才宣布发布。')
lead('页面验收','电脑和手机可打开正式网址，六项目录跳转、路线点选、地图和列表、瑞士两侧切换及属性筛选可用。来源和单位可追溯，旧值不冒充新值，结论优先趋势、数字作依据。页面不增加数据底稿板块。')

d.save(OUT)
(ROOT / 'gold-migration-demo/docs').mkdir(exist_ok=True)
shutil.copy2(OUT, ROOT / 'gold-migration-demo/docs/全球黄金迁徙地图上线需求.docx')
print(OUT)
